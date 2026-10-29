#!/usr/bin/env python3
"""Giovani Homes — read the GENERAL AVAILABILITY workbook into the data file's units[].

    python3 scripts/giovani-pricelist.py <NEW_GENERAL_AVAILABILITY_dd.mm.yy.xlsx> \
        scripts/data/giovani/2026-10-27.json scripts/data/giovani/<list date>.json

Then run scripts/giovani-sync.mjs on the output (dry run first). Needs openpyxl.

The workbook (one sheet per project, sent by email, operator decision 2026-10-28:
the price list is THE source, synced manually once a month) has quirks:
  - Hidden sheets are projects Giovani is not pushing and hidden rows are mostly
    units that are not available. Neither is a status: every row is read and the
    STATUS column decides.
  - Headers move between sheets ("COVERED VERANDA" / "COVERED VERANDAS" /
    " VERANDA", "OPEN VERANDA" / "UNCOVERED VERANDA" / " UNCOV.", Santa Elena
    has "Number"/"Bedrooms"/"Covered Areas" and TYPE twice), so columns are
    found by header text, never by position.
  - "TYPE OR FLOOR" is the house type for villas and the floor for apartments.
  - PRICELIST holds a number, 0, or the words SOLD/RESERVED. Prices are + VAT.
  - STATUS: AVAILABLE / RESERVED / SOLD map 1:1. NOT FOR SALE (landowner
    exchange units, e.g. Delear 1-3), OFF MARKET, C/N and SHOWHOUSE are not on
    the market: they are left out, so a unit we already have becomes "unlisted"
    and comes back on its own when it is offered again.
A sold unit is written without a price (same as the other manual developers).
The Giovani website feed (giovani.cy/feed/estate_property) has no status and
lists sold units, so it is not read here.
"""
import json
import re
import sys

import openpyxl

SHEETS = {  # sheet title (stripped) -> project slug in the data file
    "ALMA": "alma-villas",
    "ANGELICO": "angelico-lifestyle-apartments",
    "APANEMA": "apanema-villas",
    "AVENUE": "avenue-first",
    "CALIO": "calio",  # hidden sheet, Larnaca; units-only draft since 2026-10-29
    "CELESTIAL PERNERA": "celestial-pernera-villas",
    "DELEAR": "delear-villas",
    "ELEVATE": "elevate-premium-apartments",
    "EUPHORIA": "euphoria-lifestyle-apartments",
    "HIDDEN GEM": "hidden-gem-lifestyle-apartments",
    "NAPA AMARIS": "napa-amaris-villas",
    "NAPA MERIDIAN": "napa-meridian-villas",
    "NIERO": "niero-city-apartments",
    "PREMIERE II": "premiere-pearl-villas",
    "SABAI": "sabai-by-the-sea",
    "SANTA ELENA": "santa-elena-villas",
    "SEMELI B": "semeli-private-villas",
}
STATUS = {"AVAILABLE": "available", "RESERVED": "reserved", "SOLD": "sold"}
OFF_MARKET = {"NOT FOR SALE", "OFF MARKET", "C/N", "SHOWHOUSE"}

# header text -> field; first match wins, tested in this order
COLUMNS = [
    ("unit", r"^(UNIT NO\.?|NUMBER)$"),
    ("ptype", r"^PROPERTY TYPE$"),
    ("typefloor", r"^(TYPE OR FLOOR|FLOOR)$"),
    ("beds", r"^(BEDS|BEDROOMS)$"),
    ("internal", r"^(TOTAL COV\. AREA|COVERED AREAS)$"),
    ("veranda", r"^(COVERED VERANDAS?|VERANDA)$"),
    ("other", r"^OTHER COVERED AREAS$"),
    ("terrace", r"^(OPEN VERANDA|UNCOVERED VERANDA|UNCOV\.)$"),
    ("roof", r"ROOF GARDEN$"),
    ("basement", r"^BASEMENT$"),
    ("parking", r"PARKING$"),
    ("plot", r"^PLOT$"),
    ("pool", r"POOL$"),
    ("price", r"^PRICE"),
    ("completion", r"^COMPLETION$"),
    ("status", r"^STATUS$"),
]


def num(v):
    if isinstance(v, (int, float)):
        return v
    s = str(v or "").strip()
    return float(s) if re.fullmatch(r"\d+(\.\d+)?", s) else None


def area(v):
    n = num(v)
    return round(n, 2) if n else None


def floor_of(v):
    s = str(v if v is not None else "").strip().upper()
    if s in ("GF", "GROUND", "0"):
        return "Ground"
    m = re.match(r"^(\d+)", s)
    if m:
        return m.group(1)
    return {"FIRST": "1", "SECOND": "2", "THIRD": "3", "FOURTH": "4"}.get(s)


def read_sheet(ws):
    hdr_row = None
    for r in range(1, 15):
        vals = [re.sub(r"\s+", " ", str(c.value or "")).strip().upper() for c in ws[r]]
        if any(v.startswith("PRICE") for v in vals) and "STATUS" in vals:
            hdr_row, hdr = r, vals
            break
    if not hdr_row:
        raise SystemExit(f"{ws.title}: no header row")
    col = {}
    for field, pat in COLUMNS:
        for i, h in enumerate(hdr):
            if h and re.search(pat, h) and i not in col.values():
                col[field] = i
                break
    # Santa Elena: the first TYPE is the property type, the second the house type
    types = [i for i, h in enumerate(hdr) if h == "TYPE"]
    if types:
        col.setdefault("ptype", types[0])
        if len(types) > 1:
            col.setdefault("typefloor", types[1])
    rows = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        get = lambda f: vals[col[f]] if f in col and col[f] < len(vals) else None
        if get("unit") in (None, "") or str(get("status") or "").strip() == "":
            continue
        rows.append({f: get(f) for f in col} | {"_row": r})
    return rows


def unit_of(slug, x, problems, skipped):
    ref = re.sub(r"[\s*]", "", str(x["unit"])).upper()
    if isinstance(x["unit"], float) and x["unit"].is_integer():
        ref = str(int(x["unit"]))
    st_raw = re.sub(r"\s+", " ", str(x["status"])).strip().upper()
    if st_raw in OFF_MARKET:
        skipped.append(f"{slug} {ref} ({st_raw})")
        return None
    if st_raw not in STATUS:
        problems.append(f"{slug} {ref}: unknown status {x['status']!r} (row {x['_row']})")
        return None
    status = STATUS[st_raw]
    ptype = str(x.get("ptype") or "").strip().upper()
    if slug == "sabai-by-the-sea" and not ptype:
        ptype = "VILLA" if re.fullmatch(r"\d+", ref) else "APARTMENT"
    if "OFFICE" in ptype:
        utype = "Office"
    elif "SHOP" in ptype:
        utype = "Shop"
    elif "APARTMENT" in ptype:
        utype = "Apartment"
    elif "VILLA" in ptype or "DETACHED" in ptype or "DETTACHED" in ptype:
        utype = "Villa"
    else:
        problems.append(f"{slug} {ref}: unknown property type {x.get('ptype')!r}")
        return None
    is_flat = utype in ("Apartment", "Office", "Shop")
    beds = num(x.get("beds"))
    price = num(x.get("price"))
    u = {
        "ref": ref,
        "label": f"{utype} {ref}",
        "type": utype,
        "status": status,
        "price": int(price) if price and status != "sold" else None,
        "beds": int(beds) if beds is not None else None,
        "internal": area(x.get("internal")),
        "veranda": area(x.get("veranda")),
        "terrace": area(x.get("terrace")),
        "roof": area(x.get("roof")),
        "basement": area(x.get("basement")),
        "plot": area(x.get("plot")),
        "floor": floor_of(x.get("typefloor")) if is_flat else None,
    }
    if not is_flat and x.get("typefloor"):
        # Greek capitals sneak in (Napa Meridian "Α"); * marks the mirrored layout
        t = str(x["typefloor"]).strip().translate(str.maketrans("ΑΒΓΔΕ", "ABΓΔE"))
        u["houseType"] = f"Type {t}"
    if "DETACHED" in ptype or "DETTACHED" in ptype:
        u["houseType"] = ", ".join(filter(None, ["Semi-detached" if "SEMI" in ptype else "Detached", u.get("houseType")]))
    pool = str(x.get("pool") or "").strip().upper()
    if pool in ("PRIVATE", "COMMUNAL"):
        u["pool"] = pool.capitalize()
    other = area(x.get("other"))
    if other:
        u["otherCovered"] = other
    park = x.get("parking")
    if num(park):
        u["parking"] = f"{area(park)} m²"
    elif str(park or "").strip().upper() in ("UNCOVERED", "COVERED"):
        u["parking"] = str(park).strip().capitalize()
    if x.get("completion"):
        u["completion"] = re.sub(r"\s+", " ", str(x["completion"])).strip().title()
    if status == "available" and not u["price"]:
        problems.append(f"{slug} {ref}: AVAILABLE without a price (row {x['_row']})")
    return {k: v for k, v in u.items() if v is not None}


def main(xlsx, base, out):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    data = json.load(open(base, encoding="utf-8"))
    by_slug = {p["slug"]: p for p in data["projects"]}
    problems, skipped, seen = [], [], set()
    for ws in wb.worksheets:
        slug = SHEETS.get(ws.title.strip())
        if not slug:
            print(f"sheet {ws.title.strip()!r}: not one of our projects — ignored")
            continue
        seen.add(slug)
        units = [u for u in (unit_of(slug, x, problems, skipped) for x in read_sheet(ws)) if u]
        refs = [u["ref"] for u in units]
        dup = sorted({r for r in refs if refs.count(r) > 1})
        if dup:
            problems.append(f"{slug}: duplicate refs {dup}")
        by_slug[slug]["units"] = units
        n = {s: sum(u["status"] == s for u in units) for s in STATUS.values()}
        print(f"{slug}: {len(units)} units ({n['available']} available, {n['reserved']} reserved, {n['sold']} sold)")
    for slug in by_slug:
        if slug not in seen:
            problems.append(f"{slug}: no sheet in this workbook")
    data["_source"] = f"{data.get('_source', '')} | units: {xlsx.rsplit('/', 1)[-1]}".strip(" |")
    json.dump(data, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"\nnot on the market, left out ({len(skipped)}): {', '.join(skipped)}")
    if problems:
        print(f"\n{len(problems)} problem(s):\n  " + "\n  ".join(problems))
        sys.exit(2)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
